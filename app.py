import os, io, re, json
import streamlit as st
from dotenv import load_dotenv
from groq import Groq
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

load_dotenv()

api_key = os.getenv("GROQ_API_KEY")
client  = Groq(api_key=api_key)
MODEL   = "llama-3.3-70b-versatile"

# ══════════════════════════════════════════════════════════════
# PAGE CONFIG & CUSTOM CSS
# ══════════════════════════════════════════════════════════════
st.set_page_config(
    page_title="AI Job Agent",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');

/* ══ GLOBAL ══════════════════════════════════════════════════ */
html, body, [class*="css"] {
    font-family: 'Inter', 'Segoe UI', system-ui, sans-serif !important;
}
[data-testid="stAppViewContainer"] {
    background: linear-gradient(160deg, #eef2ff 0%, #faf9ff 45%, #f0fdf4 100%) !important;
    background-attachment: fixed !important;
}
[data-testid="stMain"] { background: transparent !important; }
[data-testid="block-container"] { padding-top: 1.5rem !important; }

/* ══ SIDEBAR ═════════════════════════════════════════════════ */
[data-testid="stSidebar"] {
    background: #ffffff !important;
    border-right: 1px solid #e5e7eb !important;
    box-shadow: 4px 0 16px rgba(79,70,229,0.06) !important;
}
[data-testid="stSidebar"] * { color: #374151 !important; }
[data-testid="stSidebar"] h1,
[data-testid="stSidebar"] h2,
[data-testid="stSidebar"] h3 { color: #111827 !important; }
[data-testid="stSidebar"] .stButton > button {
    box-shadow: 0 4px 14px rgba(79,70,229,0.30) !important;
}

/* ══ HERO ════════════════════════════════════════════════════ */
.hero {
    background: linear-gradient(135deg, #4338ca 0%, #7c3aed 55%, #a855f7 100%);
    border-radius: 20px;
    padding: 40px 46px 36px;
    margin-bottom: 32px;
    box-shadow: 0 8px 32px rgba(99,60,237,0.22), 0 2px 8px rgba(99,60,237,0.10);
    position: relative; overflow: hidden;
}
.hero::before {
    content:''; position:absolute; top:-70px; right:-50px;
    width:280px; height:280px; border-radius:50%;
    background:rgba(255,255,255,0.08); pointer-events:none;
}
.hero::after {
    content:''; position:absolute; bottom:-90px; right:80px;
    width:200px; height:200px; border-radius:50%;
    background:rgba(255,255,255,0.05); pointer-events:none;
}
.hero-badge {
    display:inline-flex; align-items:center; gap:6px;
    background:rgba(255,255,255,0.18); border:1px solid rgba(255,255,255,0.25);
    border-radius:20px; padding:5px 14px;
    font-size:0.78rem; color:rgba(255,255,255,0.95);
    font-weight:600; letter-spacing:0.02em; margin-bottom:14px;
}
.hero h1 {
    color:#ffffff !important; font-size:2.1rem !important;
    font-weight:800 !important; margin:0 0 10px 0 !important;
    letter-spacing:-0.025em !important; text-shadow:0 1px 3px rgba(0,0,0,0.15);
    position:relative; z-index:1;
}
.hero p {
    color:rgba(255,255,255,0.88) !important; font-size:1.02rem !important;
    margin:0 !important; line-height:1.65; position:relative; z-index:1;
}

/* ══ FEATURE CARDS (idle) ════════════════════════════════════ */
.feature-grid {
    display:grid; grid-template-columns:repeat(4,1fr);
    gap:16px; margin-bottom:28px;
}
.feature-card {
    background:#ffffff; border:1px solid #e5e7eb;
    border-radius:18px; padding:26px 22px; text-align:center;
    box-shadow:0 1px 4px rgba(0,0,0,0.05),0 4px 12px rgba(0,0,0,0.03);
    transition:all 0.25s ease; cursor:default;
}
.feature-card:hover {
    transform:translateY(-4px);
    box-shadow:0 8px 24px rgba(79,70,229,0.13),0 2px 8px rgba(0,0,0,0.06);
    border-color:#c4b5fd;
}
.fi-wrap {
    width:54px; height:54px; border-radius:15px;
    display:flex; align-items:center; justify-content:center;
    font-size:1.5rem; margin:0 auto 14px;
}
.fi-blue   { background:#eff6ff; }
.fi-purple { background:#f5f3ff; }
.fi-amber  { background:#fffbeb; }
.fi-green  { background:#f0fdf4; }
.feature-title { color:#111827; font-weight:700; font-size:0.93rem; margin-bottom:6px; }
.feature-desc  { color:#6b7280; font-size:0.8rem; line-height:1.55; }

/* ══ SCORE / METRIC CARDS ════════════════════════════════════ */
.card {
    background:#ffffff; border:1px solid #e5e7eb;
    border-radius:16px; padding:22px 24px; margin-bottom:12px;
    box-shadow:0 1px 3px rgba(0,0,0,0.05),0 4px 12px rgba(0,0,0,0.04);
}
.card-title {
    font-size:0.7rem; text-transform:uppercase; letter-spacing:0.1em;
    color:#9ca3af; font-weight:600; margin-bottom:10px;
}
.card-value {
    font-size:2.7rem; font-weight:800; color:#111827;
    line-height:1; letter-spacing:-0.03em;
}
.card-denom { color:#d1d5db; font-size:1.4rem; font-weight:500; }
.badge-pos {
    display:inline-flex; align-items:center;
    background:#ecfdf5; color:#059669;
    font-size:0.74rem; font-weight:700;
    padding:3px 9px; border-radius:20px; border:1px solid #a7f3d0;
    margin-left:6px; vertical-align:middle;
}
.badge-neg {
    display:inline-flex; align-items:center;
    background:#fef2f2; color:#dc2626;
    font-size:0.74rem; font-weight:700;
    padding:3px 9px; border-radius:20px; border:1px solid #fca5a5;
    margin-left:6px; vertical-align:middle;
}

/* ══ SCORE BAR ═══════════════════════════════════════════════ */
.score-wrap { margin:12px 0 4px; }
.score-bar-bg {
    background:#f3f4f6; border-radius:99px;
    height:8px; width:100%; overflow:hidden;
}
.score-bar-fill {
    height:100%; border-radius:99px;
    transition:width 0.8s cubic-bezier(0.4,0,0.2,1);
}

/* ══ KEYWORD PANEL ═══════════════════════════════════════════ */
.kw-panel {
    background:#ffffff; border:1px solid #e5e7eb;
    border-radius:18px; padding:22px 24px;
    box-shadow:0 1px 3px rgba(0,0,0,0.05),0 4px 12px rgba(0,0,0,0.04);
    margin-bottom:4px;
}
.kw-panel-header {
    display:flex; align-items:center; justify-content:space-between;
    padding-bottom:14px; margin-bottom:14px;
    border-bottom:1px solid #f3f4f6;
}
.kw-panel-label {
    font-size:0.72rem; text-transform:uppercase;
    letter-spacing:0.09em; color:#9ca3af; font-weight:600;
}
.kw-panel-stat { font-size:1.5rem; font-weight:800; line-height:1; }
.kw-section-title {
    font-size:0.72rem; font-weight:700; text-transform:uppercase;
    letter-spacing:0.08em; margin:14px 0 6px;
    display:flex; align-items:center; gap:5px;
}
.kws-matched { color:#059669; }
.kws-missing { color:#dc2626; }
.cat-label {
    font-size:0.68rem; color:#9ca3af; text-transform:uppercase;
    letter-spacing:0.08em; font-weight:600; margin:8px 0 4px;
}
.tag-matched {
    display:inline-block; background:#ecfdf5;
    border:1px solid #a7f3d0; color:#065f46;
    border-radius:20px; padding:3px 11px;
    font-size:0.75rem; font-weight:500; margin:2px 3px;
}
.tag-missing {
    display:inline-block; background:#fff1f2;
    border:1px solid #fecdd3; color:#be123c;
    border-radius:20px; padding:3px 11px;
    font-size:0.75rem; font-weight:500; margin:2px 3px;
}

/* ══ PROGRESS STEPS ══════════════════════════════════════════ */
.steps-row {
    display:flex; gap:4px; margin-bottom:24px;
    background:#f9fafb; border:1px solid #e5e7eb;
    border-radius:14px; padding:6px;
}
.step-item {
    flex:1; border-radius:10px; padding:9px 6px;
    font-size:0.73rem; text-align:center; font-weight:500;
    color:#9ca3af; background:transparent;
    white-space:nowrap; overflow:hidden; text-overflow:ellipsis;
}
.step-item.done {
    background:#ecfdf5; color:#059669; font-weight:700;
}
.step-item.active {
    background:linear-gradient(135deg,#ede9fe,#ddd6fe);
    color:#6d28d9; font-weight:700;
    box-shadow:0 1px 5px rgba(109,40,217,0.18);
}

/* ══ SECTION HEADER ══════════════════════════════════════════ */
.sec-head {
    display:flex; align-items:center; gap:12px;
    margin:4px 0 22px;
}
.sec-icon {
    width:40px; height:40px; border-radius:12px;
    display:flex; align-items:center; justify-content:center; font-size:1.15rem;
    flex-shrink:0;
}
.si-blue   { background:#eff6ff; }
.si-green  { background:#f0fdf4; }
.si-amber  { background:#fffbeb; }
.si-purple { background:#f5f3ff; }
.sec-head h2 {
    margin:0; font-size:1.25rem; font-weight:800; color:#111827;
    letter-spacing:-0.02em;
}
.sec-head p { margin:0; font-size:0.82rem; color:#6b7280; margin-top:1px; }

/* ══ DIVIDER ═════════════════════════════════════════════════ */
.divider {
    height:1px; border:none;
    background:linear-gradient(to right,transparent,#e5e7eb 15%,#e5e7eb 85%,transparent);
    margin:32px 0;
}

/* ══ SIDEBAR SPECIFICS ═══════════════════════════════════════ */
.sb-logo-wrap {
    display:flex; align-items:center; gap:11px; padding:4px 0 18px;
}
.sb-logo-icon {
    width:42px; height:42px; border-radius:12px;
    background:linear-gradient(135deg,#4f46e5,#7c3aed);
    display:flex; align-items:center; justify-content:center;
    font-size:1.25rem; flex-shrink:0;
}
.sb-logo-title { font-size:1.05rem; font-weight:800; color:#111827 !important; }
.sb-logo-sub   { font-size:0.72rem; color:#9ca3af !important; margin-top:1px; }
.sb-step-row {
    display:flex; align-items:center; gap:9px; margin-bottom:9px;
}
.sb-step-num {
    width:24px; height:24px; border-radius:50%; flex-shrink:0;
    background:linear-gradient(135deg,#4f46e5,#7c3aed);
    color:#ffffff; font-size:0.7rem; font-weight:700;
    display:flex; align-items:center; justify-content:center;
}
.sb-step-text { font-size:0.87rem; font-weight:700; color:#111827 !important; }
.hw-row {
    display:flex; align-items:flex-start; gap:8px;
    margin-bottom:7px; font-size:0.79rem; color:#6b7280 !important;
    line-height:1.5;
}
.hw-dot {
    width:6px; height:6px; border-radius:50%;
    background:#c4b5fd; flex-shrink:0; margin-top:5px;
}

/* ══ BUTTONS ═════════════════════════════════════════════════ */
[data-testid="stButton"] > button {
    background:linear-gradient(135deg,#4f46e5 0%,#7c3aed 100%) !important;
    color:#ffffff !important; border:none !important;
    border-radius:12px !important; font-weight:700 !important;
    font-size:1rem !important; padding:14px 0 !important;
    box-shadow:0 4px 14px rgba(79,70,229,0.28) !important;
    letter-spacing:0.01em !important; transition:all 0.2s ease !important;
}
[data-testid="stButton"] > button:hover {
    box-shadow:0 6px 22px rgba(79,70,229,0.38) !important;
    transform:translateY(-1px) !important;
}
[data-testid="stDownloadButton"] > button {
    background:linear-gradient(135deg,#059669 0%,#10b981 100%) !important;
    color:#ffffff !important; border:none !important;
    border-radius:12px !important; font-weight:700 !important;
    font-size:1.05rem !important; padding:16px 0 !important;
    box-shadow:0 4px 14px rgba(16,185,129,0.28) !important;
    letter-spacing:0.01em !important; transition:all 0.2s ease !important;
}
[data-testid="stDownloadButton"] > button:hover {
    box-shadow:0 6px 22px rgba(16,185,129,0.38) !important;
    transform:translateY(-1px) !important;
}

/* ══ INPUTS ══════════════════════════════════════════════════ */
textarea {
    background:#fafafa !important; color:#111827 !important;
    border:1.5px solid #e5e7eb !important; border-radius:10px !important;
    font-family:'Inter','Segoe UI',sans-serif !important;
    font-size:0.87rem !important; line-height:1.65 !important;
    transition:border-color 0.2s, box-shadow 0.2s !important;
}
textarea:focus {
    border-color:#a5b4fc !important;
    box-shadow:0 0 0 3px rgba(165,180,252,0.22) !important;
}
[data-testid="stFileUploader"] {
    background:#fafafa !important;
    border:2px dashed #d1d5db !important;
    border-radius:14px !important; padding:10px !important;
}

/* ══ ALERTS ══════════════════════════════════════════════════ */
[data-testid="stAlert"] {
    border-radius:12px !important; border-left-width:4px !important;
}

/* ══ TABS ════════════════════════════════════════════════════ */
[data-testid="stTabs"] [data-baseweb="tab-list"] {
    background:#f3f4f6 !important; border-radius:12px !important;
    padding:4px !important; gap:4px !important;
}
[data-testid="stTabs"] [data-baseweb="tab"] {
    border-radius:9px !important; color:#6b7280 !important; font-weight:500 !important;
}
[data-testid="stTabs"] [aria-selected="true"] {
    background:#ffffff !important; color:#4f46e5 !important;
    box-shadow:0 1px 5px rgba(0,0,0,0.1) !important; font-weight:700 !important;
}

/* ══ CODE BLOCKS ══════════════════════════════════════════════ */
[data-testid="stCode"] {
    background:#f8fafc !important; border:1px solid #e2e8f0 !important;
    border-radius:10px !important;
}

/* ══ SPINNER ═════════════════════════════════════════════════ */
[data-testid="stSpinner"] { color:#4f46e5 !important; }

/* ══ HIDE BRANDING ══════════════════════════════════════════ */
#MainMenu, footer { visibility:hidden; }
</style>
""", unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════
# AI CALL
# ══════════════════════════════════════════════════════════════
def call_ai(system_prompt, user_prompt, temperature=0.3):
    try:
        r = client.chat.completions.create(
            model=MODEL,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user",   "content": user_prompt},
            ],
            temperature=temperature,
        )
        return r.choices[0].message.content.replace("|", " ")
    except Exception as e:
        st.error(f"AI call failed: {e}")
        return ""


# ══════════════════════════════════════════════════════════════
# STEP 1 — Extract real keywords from the JD
# ══════════════════════════════════════════════════════════════
def extract_jd_keywords(jd: str) -> dict:
    system = (
        "You are an expert ATS keyword analyst. "
        "Read the job description carefully and extract every skill, tool, technology, "
        "qualification, and soft skill the employer requires or prefers. "
        "Be thorough — include both hard skills and soft skills. "
        "Group them into logical categories (e.g. Technical Skills, Tools, Soft Skills, etc.). "
        "Return ONLY valid JSON, no markdown, no explanation, no backticks. "
        "Format: {\"Category\": [\"keyword1\", \"keyword2\"], ...} "
        "All keywords lowercase."
    )
    raw = call_ai(system, f"JOB DESCRIPTION:\n{jd}", temperature=0.1)
    raw = re.sub(r"```[a-z]*", "", raw).strip().strip("`").strip()
    try:
        data = json.loads(raw)
        return {cat: [k.lower().strip() for k in kws]
                for cat, kws in data.items() if isinstance(kws, list)}
    except Exception:
        return {}


# ══════════════════════════════════════════════════════════════
# STEP 2 — Honest ATS scoring against real JD keywords
# ══════════════════════════════════════════════════════════════
def score_cv(cv_text: str, jd_keywords: dict):
    cv_lower = cv_text.lower()
    matched_by_cat, missing_by_cat = {}, {}
    total_jd = total_matched = 0

    for cat, keywords in jd_keywords.items():
        m    = [kw for kw in keywords if kw in cv_lower]
        miss = [kw for kw in keywords if kw not in cv_lower]
        if m:    matched_by_cat[cat] = m
        if miss: missing_by_cat[cat] = miss
        total_jd      += len(keywords)
        total_matched += len(m)

    kw_score = int((total_matched / max(total_jd, 1)) * 60)

    exp = 0
    if any(w in cv_lower for w in ["experience", "work experience"]): exp += 8
    if any(w in cv_lower for w in ["developed","created","built","analyzed",
           "optimized","automated","implemented","managed","resolved",
           "configured","deployed","supported","maintained","improved",
           "reduced","led","delivered"]):                              exp += 9
    if any(w in cv_lower for w in ["bachelor","master","bsc","msc","degree"]): exp += 8

    fmt = 15
    if len(cv_text.split()) > 900:       fmt -= 4
    if "|" in cv_text:                   fmt -= 3
    if "skills" not in cv_lower:         fmt -= 4
    if "experience" not in cv_lower:     fmt -= 4

    return min(100, kw_score + exp + max(fmt, 0)), matched_by_cat, missing_by_cat


# ══════════════════════════════════════════════════════════════
# STEP 3 — Read DOCX text (paragraphs + tables)
# ══════════════════════════════════════════════════════════════
def read_docx(file) -> str:
    try:
        file.seek(0)
        doc = Document(file)
        lines = [p.text for p in doc.paragraphs]
        for tbl in doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p.text.strip():
                            lines.append(p.text)
        return "\n".join(lines)
    except Exception as e:
        st.error(f"Failed to read CV: {e}")
        return ""


# ══════════════════════════════════════════════════════════════
# STEP 4 — Rebuild CV  (skills: keep + enrich + prune)
# ══════════════════════════════════════════════════════════════
def get_section(text, name):
    pattern = rf"(?m)^{name}:\s*(.*?)(?=^[A-Z][A-Z_/ ]+:|\Z)"
    m = re.search(pattern, text, re.DOTALL)
    return m.group(1).strip().replace("|", " ") if m else ""


def rebuild_cv(old_cv: str, jd: str, jd_keywords: dict) -> dict:
    flat_kw = ", ".join(kw for kws in jd_keywords.values() for kw in kws)

    prompt = f"""
You are rebuilding a professional ATS-optimized CV for a new job.

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
STRICT RULES
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
1. REAL FACTS ONLY — keep all real companies, dates, universities, degrees, name, contact.
2. WORK EXPERIENCE — rewrite every bullet to focus on the target job. Use strong action verbs.
   Weave these JD keywords naturally into the bullets: {flat_kw}
3. SKILLS — this is the most important section. Follow these sub-rules exactly:
   a. READ every skill in the OLD CV carefully.
   b. KEEP any skill that is relevant OR transferable to the target job (even partially).
   c. REMOVE only skills that are completely unrelated to the target job (e.g. if target is
      IT Support, remove "DAX formulas" but keep "Python scripting").
   d. ADD every skill from the JD keywords list that is not already present.
   e. ADD any additional skills you know are commonly expected for this role but not listed.
   f. Organise all skills into clean categories. The goal is MAXIMUM relevant skills.
4. CERTIFICATIONS — keep all certs from the old CV; add JD-relevant ones only if real.
5. No invented experience, no fake tools, no vertical bar symbol.
6. Output must be concise and one-page ready.

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
OUTPUT FORMAT — return exactly these labels, NOTHING before TITLE:
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

TITLE:
<exact target job title>

SKILLS:
<Category 1>: <skill>, <skill>, <skill>, ...
<Category 2>: <skill>, <skill>, <skill>, ...
<Category 3>: <skill>, <skill>, <skill>, ...
Soft Skills: <skill>, <skill>, <skill>

WORK_EXPERIENCE:
<Job Title> (<Company>) >>> <Date Range>, <Country>
- <strong action verb> + <achievement> + <JD keyword>
- <bullet>
- <bullet>
- <bullet>
- <bullet>

<Job Title> - <Company> >>> <Date Range>, <Country>
- <bullet>
- <bullet>
- <bullet>

EDUCATION:
<University Name> >>> <Date Range>
<Degree Title> >>> <City>, <Country>

<University Name> >>> <Date Range>
<Degree Title> >>> <City>, <Country>

CERTIFICATION:
- <cert name>
- <cert name>

LANGUAGE:
- English, Fluent
- German, Beginner

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
OLD CV (use only for real facts):
{old_cv}

JOB DESCRIPTION (target role):
{jd}
"""
    system = (
        "You are an expert ATS CV writer. "
        "Preserve and enrich skills from the old CV. "
        "Never invent experience. Never use vertical bars. "
        "Return only the formatted CV, starting with TITLE:"
    )
    raw = call_ai(system, prompt, temperature=0.3)

    match = re.search(r"(?m)^TITLE:", raw)
    if match:
        raw = raw[match.start():]

    return {
        "raw":           raw,
        "title":         get_section(raw, "TITLE"),
        "skills":        get_section(raw, "SKILLS"),
        "work":          get_section(raw, "WORK_EXPERIENCE"),
        "education":     get_section(raw, "EDUCATION"),
        "certification": get_section(raw, "CERTIFICATION"),
        "language":      get_section(raw, "LANGUAGE"),
    }


# ══════════════════════════════════════════════════════════════
# STEP 5 — Build DOCX using pure python-docx (no Node.js needed)
# ══════════════════════════════════════════════════════════════
_MON = (r'(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?'
        r'|Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)')
_DATE_RANGE_RE = re.compile(
    rf'(?:{_MON}\.?\s+)?\d{{4}}\s*[-–—]\s*(?:(?:{_MON}\.?\s+)?\d{{4}}|Present|present|Current)',
    re.IGNORECASE,
)
_TRAILING_MON_RE = re.compile(rf'({_MON}\.?\s+)$', re.IGNORECASE)

def _split_left_right(line: str):
    if '>>>' in line:
        a, b = line.split('>>>', 1)
        return a.strip(), b.strip()
    m = _DATE_RANGE_RE.search(line)
    if m:
        pre   = line[:m.start()]
        mon_m = _TRAILING_MON_RE.search(pre)
        cut   = m.start() - len(mon_m.group(1)) if mon_m else m.start()
        return line[:cut].rstrip(' ,'), line[cut:].strip()
    idx = line.rfind(',')
    if idx > 0:
        return line[:idx].strip(), line[idx + 1:].strip()
    return line.strip(), ''


def _parse_edu_pairs(edu_text: str):
    INST = {'University','Institute','College','School','Technology',
            'Academy','Fachhochschule','Hochschule'}
    lines = [l.strip() for l in edu_text.split('\n') if l.strip()]
    entries = []
    i = 0
    while i < len(lines):
        raw = lines[i]
        if '>>>' in raw:
            left, right = [p.strip() for p in raw.split('>>>', 1)]
            is_inst = any(w in left for w in INST)
            if is_inst:
                uni, date = left, right
                degree = location = ''
                if i + 1 < len(lines):
                    nxt = lines[i + 1]
                    if '>>>' in nxt:
                        d, loc = nxt.split('>>>', 1)
                        degree, location = d.strip(), loc.strip()
                    else:
                        degree = nxt.strip()
                    i += 1
                entries.append((uni, date, degree, location))
            else:
                entries.append(('', '', left, right))
        else:
            is_inst = any(w in raw for w in INST)
            if is_inst:
                m = _DATE_RANGE_RE.search(raw)
                if m:
                    pre   = raw[:m.start()]
                    mon_m = _TRAILING_MON_RE.search(pre)
                    cut   = m.start() - len(mon_m.group(1)) if mon_m else m.start()
                    uni      = raw[:cut].rstrip(' ,')
                    date     = raw[cut:m.end()].strip()
                    location = raw[m.end():].lstrip(' ,').strip()
                else:
                    parts    = [p.strip() for p in raw.split(',')]
                    uni      = parts[0]
                    date     = parts[1] if len(parts) > 1 else ''
                    location = ', '.join(parts[2:]) if len(parts) > 2 else ''
                degree = ''
                if i + 1 < len(lines):
                    nxt = lines[i + 1]
                    nxt_is_inst = any(w in nxt.replace('>>>', '') for w in INST)
                    if not nxt_is_inst:
                        if '>>>' in nxt:
                            d, loc = nxt.split('>>>', 1)
                            degree = d.strip()
                            if not location:
                                location = loc.strip()
                        else:
                            degree = nxt.strip()
                        i += 1
                entries.append((uni, date, degree, location))
            else:
                if '>>>' in raw:
                    d, loc = raw.split('>>>', 1)
                    entries.append(('', '', d.strip(), loc.strip()))
                else:
                    entries.append(('', '', raw, ''))
        i += 1
    return entries


# ── python-docx helpers ───────────────────────────────────────
_RIGHT_TAB_POS = Inches(7.8)   # content width: 8.5" - 0.35" - 0.35"

def _add_bottom_border(para):
    pPr  = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    bot.set(qn('w:val'),   'single')
    bot.set(qn('w:sz'),    '12')
    bot.set(qn('w:space'), '1')
    bot.set(qn('w:color'), '000000')
    pBdr.append(bot)
    pPr.append(pBdr)

def _add_right_tab(para):
    pPr    = para._p.get_or_add_pPr()
    tabs   = OxmlElement('w:tabs')
    tab    = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:pos'), str(int(7.8 * 1440)))   # 11232 twips
    tabs.append(tab)
    pPr.append(tabs)

def _sp(para, before=0, after=0, line=1.15):
    para.paragraph_format.space_before  = Pt(before)
    para.paragraph_format.space_after   = Pt(after)
    para.paragraph_format.line_spacing  = line

def _run(para, text, bold=False, size=10):
    r = para.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(size)
    r.bold      = bold
    return r


def create_docx(data: dict):
    """Build the CV DOCX entirely with python-docx — works on Streamlit Cloud."""
    doc = Document()

    # ── Page margins ──────────────────────────────────────────
    sec = doc.sections[0]
    sec.top_margin    = Inches(0.35)
    sec.bottom_margin = Inches(0.30)
    sec.left_margin   = Inches(0.35)
    sec.right_margin  = Inches(0.35)

    # ── Default font ──────────────────────────────────────────
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(10)

    # ── Section heading helper ────────────────────────────────
    def heading(title):
        p = doc.add_paragraph()
        _sp(p, before=7, after=3)
        _add_bottom_border(p)
        _run(p, title, bold=True, size=12)

    # ── NAME + LOCATION ───────────────────────────────────────
    p = doc.add_paragraph()
    _sp(p, before=0, after=0)
    _add_right_tab(p)
    _run(p, 'Junaid Hasan', bold=True, size=17)
    _run(p, '\tGermany', size=11)

    # ── TITLE + CONTACT ───────────────────────────────────────
    p2 = doc.add_paragraph()
    _sp(p2, before=0, after=4)
    _add_right_tab(p2)
    _add_bottom_border(p2)
    _run(p2, data['title'], size=11)
    _run(p2, '\tTel: +4915563388607   E-mail: junaidhasan696@gmail.com   LinkedIn', size=10)

    # ── SKILLS ────────────────────────────────────────────────
    heading('SKILLS')
    for line in data['skills'].split('\n'):
        line = line.strip()
        if not line: continue
        p = doc.add_paragraph()
        _sp(p, after=1)
        if ':' in line:
            cat, rest = line.split(':', 1)
            _run(p, cat.strip() + ': ', bold=True)
            _run(p, rest.strip())
        else:
            _run(p, line)

    # ── WORK EXPERIENCE ───────────────────────────────────────
    heading('WORK EXPERIENCE')
    for line in data['work'].split('\n'):
        line = line.strip()
        if not line: continue
        if line.startswith('-') or line.startswith('•'):
            bullet = line.lstrip('-•· ').strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent       = Inches(0.20)
            p.paragraph_format.first_line_indent = Inches(-0.15)
            _sp(p, after=0)
            _run(p, '•  ' + bullet)
        else:
            left, right = _split_left_right(line)
            p = doc.add_paragraph()
            _sp(p, before=4, after=1)
            if right:
                _add_right_tab(p)
                _run(p, left, bold=True)
                _run(p, '\t' + right)
            else:
                _run(p, left, bold=True)

    # ── EDUCATION ─────────────────────────────────────────────
    heading('EDUCATION')
    for (uni, date, degree, location) in _parse_edu_pairs(data['education']):
        if uni:
            p = doc.add_paragraph()
            _sp(p, before=3, after=0)
            _add_right_tab(p)
            _run(p, uni, bold=True)
            _run(p, '\t' + date)
        if degree or location:
            p = doc.add_paragraph()
            _sp(p, before=0, after=1)
            _add_right_tab(p)
            _run(p, degree)
            _run(p, '\t' + location)

    # ── CERTIFICATION ─────────────────────────────────────────
    heading('CERTIFICATION')
    for line in data['certification'].split('\n'):
        line = line.strip().lstrip('-• ').strip()
        if not line: continue
        p = doc.add_paragraph()
        _sp(p, after=1)
        _run(p, line)

    # ── LANGUAGE ──────────────────────────────────────────────
    heading('Language')
    for line in data['language'].split('\n'):
        line = line.strip().lstrip('-• ').strip()
        if not line: continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent       = Inches(0.20)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        _sp(p, after=0)
        _run(p, '•  ' + line)

    # ── Save to buffer ────────────────────────────────────────
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ══════════════════════════════════════════════════════════════
# UI HELPERS
# ══════════════════════════════════════════════════════════════
def score_color(score: int) -> str:
    if score >= 80: return "#059669"
    if score >= 55: return "#d97706"
    return "#dc2626"

def score_gradient(score: int) -> str:
    if score >= 80: return "linear-gradient(90deg,#10b981,#34d399)"
    if score >= 55: return "linear-gradient(90deg,#f59e0b,#fbbf24)"
    return "linear-gradient(90deg,#ef4444,#f87171)"

def score_bar(score: int, _color: str):
    grad = score_gradient(score)
    st.markdown(f"""
    <div class="score-wrap">
      <div class="score-bar-bg">
        <div class="score-bar-fill" style="width:{score}%;background:{grad};"></div>
      </div>
    </div>""", unsafe_allow_html=True)


def keyword_panel(label: str, matched: dict, missing: dict, total: int):
    n_match = sum(len(v) for v in matched.values())
    pct     = int(n_match / max(total, 1) * 100)
    color   = score_color(pct)

    st.markdown(f"""
    <div class="kw-panel">
      <div class="kw-panel-header">
        <span class="kw-panel-label">{label}</span>
        <span class="kw-panel-stat" style="color:{color};">
          {n_match}<span style="color:#d1d5db;font-size:1rem;font-weight:400;">/{total}</span>
          &nbsp;<span style="font-size:0.82rem;color:{color};font-weight:700;">{pct}%</span>
        </span>
      </div>
    </div>""", unsafe_allow_html=True)

    if matched:
        st.markdown('<div class="kw-section-title kws-matched">✅ Matched</div>',
                    unsafe_allow_html=True)
        for cat, kws in matched.items():
            tags = "".join(f'<span class="tag-matched">{k}</span>' for k in kws)
            st.markdown(f'<div class="cat-label">{cat}</div><div style="margin-bottom:4px;">{tags}</div>',
                        unsafe_allow_html=True)

    if missing:
        st.markdown('<div class="kw-section-title kws-missing" style="margin-top:16px;">❌ Missing</div>',
                    unsafe_allow_html=True)
        for cat, kws in missing.items():
            tags = "".join(f'<span class="tag-missing">{k}</span>' for k in kws)
            st.markdown(f'<div class="cat-label">{cat}</div><div style="margin-bottom:4px;">{tags}</div>',
                        unsafe_allow_html=True)
    elif n_match > 0:
        st.success("🎉 All JD keywords present!")


# ══════════════════════════════════════════════════════════════
# SIDEBAR
# ══════════════════════════════════════════════════════════════
with st.sidebar:
    st.markdown("""
    <div class="sb-logo-wrap">
      <div class="sb-logo-icon">🎯</div>
      <div>
        <div class="sb-logo-title">AI Job Agent</div>
        <div class="sb-logo-sub">ATS CV Optimizer</div>
      </div>
    </div>
    <hr style="border:none;border-top:1px solid #e5e7eb;margin:0 0 22px;">
    """, unsafe_allow_html=True)

    st.markdown("""
    <div class="sb-step-row">
      <div class="sb-step-num">1</div>
      <div class="sb-step-text">Upload your CV</div>
    </div>
    """, unsafe_allow_html=True)
    old_cv_file = st.file_uploader(
        "Upload your current CV", type=["docx"], label_visibility="collapsed",
    )
    if old_cv_file:
        st.success(f"✅ {old_cv_file.name}")

    st.markdown("<div style='height:22px'></div>", unsafe_allow_html=True)

    st.markdown("""
    <div class="sb-step-row">
      <div class="sb-step-num">2</div>
      <div class="sb-step-text">Paste Job Description</div>
    </div>
    """, unsafe_allow_html=True)
    job_desc = st.text_area(
        "Job description", height=295,
        placeholder="Paste the full job description here…",
        label_visibility="collapsed",
    )
    if job_desc.strip():
        st.caption(f"📝 {len(job_desc.split())} words")

    st.markdown("<div style='height:18px'></div>", unsafe_allow_html=True)
    go = st.button("🚀 Rebuild My CV", use_container_width=True)

    st.markdown("<hr style='border:none;border-top:1px solid #e5e7eb;margin:26px 0 16px;'>",
                unsafe_allow_html=True)
    st.markdown("""
    <div style="font-size:0.72rem;font-weight:700;text-transform:uppercase;
                letter-spacing:0.09em;color:#9ca3af;margin-bottom:12px;">
      How it works
    </div>
    """, unsafe_allow_html=True)
    for step in [
        "Keywords extracted from <i>your actual JD</i>",
        "Old CV scored honestly against them",
        "Skills kept, enriched &amp; JD-aligned",
        "Experience rewritten with JD keywords",
        "New CV scored &amp; DOCX generated",
    ]:
        st.markdown(f'<div class="hw-row"><div class="hw-dot"></div><div>{step}</div></div>',
                    unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════
# MAIN CONTENT
# ══════════════════════════════════════════════════════════════
st.markdown("""
<div class="hero">
  <div style="position:relative;z-index:1;">
    <div class="hero-badge">✨ AI-Powered ATS Optimization</div>
    <h1>Job Application Assistant</h1>
    <p>Transform any CV for any job in seconds — real keyword analysis,<br>
       honest scoring, and a download-ready Word file every time.</p>
  </div>
</div>
""", unsafe_allow_html=True)

# ── IDLE STATE ─────────────────────────────────────────────────────────────
if not go:
    st.markdown("""
    <div class="feature-grid">
      <div class="feature-card">
        <div class="fi-wrap fi-blue">🔍</div>
        <div class="feature-title">Real JD Keywords</div>
        <div class="feature-desc">AI reads your actual job description, not a fixed generic list</div>
      </div>
      <div class="feature-card">
        <div class="fi-wrap fi-purple">🧠</div>
        <div class="feature-title">Skills Preserved</div>
        <div class="feature-desc">Keeps relevant skills from your old CV, prunes only irrelevant ones</div>
      </div>
      <div class="feature-card">
        <div class="fi-wrap fi-amber">📊</div>
        <div class="feature-title">Honest Scoring</div>
        <div class="feature-desc">Score reflects true keyword match against the specific JD only</div>
      </div>
      <div class="feature-card">
        <div class="fi-wrap fi-green">📄</div>
        <div class="feature-title">Word File Ready</div>
        <div class="feature-desc">Styled DOCX matching your original CV layout and formatting</div>
      </div>
    </div>
    """, unsafe_allow_html=True)

    st.markdown("""
    <div style="background:#ffffff;border:1px solid #e5e7eb;border-left:4px solid #6366f1;
                border-radius:12px;padding:16px 20px;display:flex;align-items:center;gap:12px;
                box-shadow:0 1px 4px rgba(0,0,0,0.04);">
      <span style="font-size:1.4rem;">👈</span>
      <span style="color:#374151;font-size:0.93rem;">
        Upload your CV and paste a job description in the sidebar,
        then click <strong style="color:#4f46e5;">Rebuild My CV</strong> to get started.
      </span>
    </div>
    """, unsafe_allow_html=True)
    st.stop()

# ── VALIDATION ──────────────────────────────────────────────────────────────
if not old_cv_file:
    st.error("Please upload your CV in the sidebar."); st.stop()
if not job_desc.strip():
    st.error("Please paste the job description in the sidebar."); st.stop()
if not api_key:
    st.error("GROQ_API_KEY not found in .env file."); st.stop()

old_text = read_docx(old_cv_file)
if not old_text.strip():
    st.error("Could not read CV text — please check the file."); st.stop()

# ── PROGRESS ────────────────────────────────────────────────────────────────
progress_ph = st.empty()
STEPS = ["🔍 Extract Keywords", "📋 Score Before", "✍️ Rebuild CV",
         "📈 Score After", "📝 Build DOCX"]

def show_progress(done, active, all_steps):
    items = ""
    for s in all_steps:
        cls    = "done" if s in done else ("active" if s == active else "step-item")
        prefix = "✓ "  if s in done else ("▶ "      if s == active else "")
        items += f'<div class="step-item {cls}">{prefix}{s}</div>'
    progress_ph.markdown(f'<div class="steps-row">{items}</div>', unsafe_allow_html=True)

show_progress([], STEPS[0], STEPS)
with st.spinner("Extracting keywords from job description…"):
    jd_keywords = extract_jd_keywords(job_desc)

if not jd_keywords:
    st.warning("Could not extract keywords — check your API key."); st.stop()

total_jd_kw = sum(len(v) for v in jd_keywords.values())

show_progress(STEPS[:1], STEPS[1], STEPS)
before_score, before_matched, before_missing = score_cv(old_text, jd_keywords)

show_progress(STEPS[:2], STEPS[2], STEPS)
with st.spinner("Rebuilding CV — preserving your skills, adding JD keywords…"):
    data = rebuild_cv(old_text, job_desc, jd_keywords)

if not data["raw"].strip():
    st.error("CV generation failed."); st.stop()

new_text = "\n".join([data["title"], data["skills"], data["work"],
                      data["education"], data["certification"], data["language"]])

show_progress(STEPS[:3], STEPS[3], STEPS)
after_score, after_matched, after_missing = score_cv(new_text, jd_keywords)

show_progress(STEPS[:4], STEPS[4], STEPS)
with st.spinner("Generating Word document…"):
    docx_buf = create_docx(data)

show_progress(STEPS, "", STEPS)

st.markdown("""
<div style="background:linear-gradient(135deg,#ecfdf5,#d1fae5);
            border:1px solid #6ee7b7;border-radius:14px;
            padding:16px 22px;display:flex;align-items:center;gap:12px;
            margin:8px 0 24px;box-shadow:0 2px 8px rgba(16,185,129,0.12);">
  <span style="font-size:1.5rem;">🎉</span>
  <div>
    <div style="font-weight:700;color:#065f46;font-size:0.95rem;">CV successfully rebuilt!</div>
    <div style="font-size:0.82rem;color:#047857;margin-top:2px;">
      Your ATS-optimized CV is ready to download below.
    </div>
  </div>
</div>
""", unsafe_allow_html=True)

# ══════════════════════════════════════════════════════════════
# RESULTS
# ══════════════════════════════════════════════════════════════

# ── ATS SCORE ───────────────────────────────────────────────────────────────
st.markdown("""
<div class="sec-head">
  <div class="sec-icon si-amber">📊</div>
  <div><h2>ATS Score Comparison</h2><p>Scored against keywords extracted from this job description only</p></div>
</div>""", unsafe_allow_html=True)

bm          = sum(len(v) for v in before_matched.values())
am          = sum(len(v) for v in after_matched.values())
delta_score = after_score - before_score
delta_kw    = am - bm

c1, c2, c3, c4 = st.columns(4)

with c1:
    bc = score_color(before_score)
    st.markdown(f"""<div class="card">
      <div class="card-title">Before Score</div>
      <div class="card-value" style="color:{bc};">{before_score}<span class="card-denom">/100</span></div>
    </div>""", unsafe_allow_html=True)
    score_bar(before_score, bc)

with c2:
    ac   = score_color(after_score)
    delt = (f'<span class="badge-pos">▲ +{delta_score}</span>' if delta_score > 0
            else f'<span class="badge-neg">▼ {delta_score}</span>')
    st.markdown(f"""<div class="card">
      <div class="card-title">After Score {delt}</div>
      <div class="card-value" style="color:{ac};">{after_score}<span class="card-denom">/100</span></div>
    </div>""", unsafe_allow_html=True)
    score_bar(after_score, ac)

with c3:
    st.markdown(f"""<div class="card">
      <div class="card-title">Keywords Before</div>
      <div class="card-value">{bm}<span class="card-denom">/{total_jd_kw}</span></div>
    </div>""", unsafe_allow_html=True)

with c4:
    kd = (f'<span class="badge-pos">▲ +{delta_kw}</span>' if delta_kw > 0
          else f'<span class="badge-neg">▼ {delta_kw}</span>')
    st.markdown(f"""<div class="card">
      <div class="card-title">Keywords After {kd}</div>
      <div class="card-value">{am}<span class="card-denom">/{total_jd_kw}</span></div>
    </div>""", unsafe_allow_html=True)

# ── KEYWORD ANALYSIS ────────────────────────────────────────────────────────
st.markdown("<hr class='divider'>", unsafe_allow_html=True)
st.markdown(f"""
<div class="sec-head">
  <div class="sec-icon si-blue">🔑</div>
  <div><h2>Keyword Analysis</h2>
       <p>{total_jd_kw} keywords across {len(jd_keywords)} categories extracted from this JD</p></div>
</div>""", unsafe_allow_html=True)

ka, kb = st.columns(2)
with ka:
    keyword_panel("Before — Old CV vs JD", before_matched, before_missing, total_jd_kw)
with kb:
    keyword_panel("After — Rebuilt CV vs JD", after_matched, after_missing, total_jd_kw)

# ── CV PREVIEW ───────────────────────────────────────────────────────────────
st.markdown("<hr class='divider'>", unsafe_allow_html=True)
st.markdown("""
<div class="sec-head">
  <div class="sec-icon si-purple">📄</div>
  <div><h2>Rebuilt CV Preview</h2><p>Review your new CV content before downloading</p></div>
</div>""", unsafe_allow_html=True)

tab1, tab2 = st.tabs(["📝 Full Text", "🔧 Section by Section"])
with tab1:
    st.text_area("Rebuilt CV", new_text, height=440, label_visibility="collapsed")

with tab2:
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("**🏷️ Title**")
        st.code(data["title"], language=None)
        st.markdown("**🎓 Education**")
        st.code(data["education"], language=None)
        st.markdown("**📜 Certifications**")
        st.code(data["certification"], language=None)
    with c2:
        st.markdown("**🛠️ Skills**")
        st.code(data["skills"], language=None)
        st.markdown("**🌍 Languages**")
        st.code(data["language"], language=None)
    st.markdown("**💼 Work Experience**")
    st.code(data["work"], language=None)

# ── DOWNLOAD ─────────────────────────────────────────────────────────────────
st.markdown("<hr class='divider'>", unsafe_allow_html=True)
if docx_buf:
    st.markdown("""
    <div style="text-align:center;margin-bottom:14px;">
      <div style="font-size:0.78rem;text-transform:uppercase;letter-spacing:0.09em;
                  color:#9ca3af;font-weight:600;margin-bottom:6px;">Ready to apply?</div>
      <div style="font-size:1.05rem;font-weight:700;color:#111827;margin-bottom:16px;">
        Download your rebuilt, ATS-optimized CV
      </div>
    </div>
    """, unsafe_allow_html=True)
    st.download_button(
        "⬇️  Download Rebuilt CV (.docx)",
        data=docx_buf,
        file_name="rebuilt_ats_cv.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True,
    )
else:
    st.warning("DOCX could not be generated — download unavailable.")